"""Dataset rules: who can see and change what, versions, imports, pipelines, quarantine,
sharing, export and promotion to the catalogue. All data work goes through the engine package,
so no SQL or code ever comes from a client."""

from __future__ import annotations

import csv
import io
import uuid
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from pydantic import TypeAdapter
from sqlalchemy import ColumnElement, Select, exists, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import Conflict, Forbidden, NotFound, StaleVersion, ValidationFailed
from app.core.timeutil import utcnow
from app.modules.audit_log.contracts import record
from app.modules.catalogue import contracts as catalogue
from app.modules.datasets import storage
from app.modules.datasets.engine import charts, eda, frame, infer
from app.modules.datasets.engine import compile as pipeline
from app.modules.datasets.engine.steps import Join, Step
from app.modules.datasets.engine.validate import Rule, apply_rules
from app.modules.datasets.models import (
    Dataset,
    DatasetShare,
    DatasetVersion,
    LibraryFile,
    PromotionItem,
    QuarantineRow,
    Synonym,
)
from app.modules.files.contracts import read_file
from app.modules.identity.contracts import P, Principal, audit_context, ensure_different_people

_STEPS = TypeAdapter(list[Step])
_RULES = TypeAdapter(list[Rule])
MAX_PROMOTE = 500
EXPORT_ROWS_XLSX = 200_000


# ------------------------------------------------------------------ access


def _shared_with(principal: Principal, access: tuple[str, ...]) -> ColumnElement[bool]:
    roles = [r.value for r in principal.roles]
    return exists().where(
        DatasetShare.dataset_id == Dataset.id,
        DatasetShare.access.in_(access),
        or_(
            (DatasetShare.kind == "user") & (DatasetShare.principal == str(principal.user_id)),
            (DatasetShare.kind == "role") & DatasetShare.principal.in_(roles or [""]),
        ),
    )


def visible(principal: Principal) -> ColumnElement[bool]:
    if principal.has(P.DATASET_ADMIN):
        return exists().where(Dataset.id == Dataset.id)
    return or_(Dataset.owner_id == principal.user_id, _shared_with(principal, ("view", "edit")))


def datasets_query(
    principal: Principal, *, q: str | None, tag: str | None, include_deleted: bool
) -> Select[Dataset]:
    stmt = select(Dataset).where(visible(principal)).order_by(Dataset.name)
    if not include_deleted:
        stmt = stmt.where(Dataset.deleted_at.is_(None))
    if q:
        stmt = stmt.where(Dataset.name.ilike(f"%{q}%"))
    if tag:
        stmt = stmt.where(Dataset.tags.contains([tag]))
    return stmt


async def get_dataset(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    *,
    edit: bool = False,
    deleted: bool = False,
) -> Dataset:
    principal.require(P.DATASET_READ)
    ds = await session.scalar(select(Dataset).where(Dataset.id == dataset_id, visible(principal)))
    if ds is None or (ds.deleted_at is not None and not deleted):
        raise NotFound("Dataset not found.")
    if edit:
        principal.require(P.DATASET_WRITE)
        can = (
            ds.owner_id == principal.user_id
            or principal.has(P.DATASET_ADMIN)
            or bool(
                await session.scalar(
                    select(
                        exists().where(
                            DatasetShare.dataset_id == ds.id,
                            DatasetShare.access == "edit",
                            _edit_match(principal),
                        )
                    )
                )
            )
        )
        if not can:
            raise Forbidden("This dataset is shared with you to view only.")
    return ds


def _edit_match(principal: Principal) -> ColumnElement[bool]:
    roles = [r.value for r in principal.roles]
    return or_(
        (DatasetShare.kind == "user") & (DatasetShare.principal == str(principal.user_id)),
        (DatasetShare.kind == "role") & DatasetShare.principal.in_(roles or [""]),
    )


def _owner_or_admin(principal: Principal, ds: Dataset) -> None:
    if ds.owner_id != principal.user_id and not principal.has(P.DATASET_ADMIN):
        raise Forbidden("Only the owner or an admin can do this.")


# ------------------------------------------------------------------ versions


async def get_version(
    session: AsyncSession, ds: Dataset, number: int | None = None
) -> DatasetVersion:
    stmt = select(DatasetVersion).where(DatasetVersion.dataset_id == ds.id)
    stmt = stmt.where(DatasetVersion.number == (number or ds.latest_version))
    v = await session.scalar(stmt)
    if v is None:
        raise NotFound("That version does not exist.")
    return v


async def list_versions(session: AsyncSession, ds: Dataset) -> list[DatasetVersion]:
    rows = await session.scalars(
        select(DatasetVersion)
        .where(DatasetVersion.dataset_id == ds.id)
        .order_by(DatasetVersion.number.desc())
    )
    return list(rows)


async def load_table(v: DatasetVersion) -> pa.Table:
    return await storage.get_table(v.storage_key)


async def write_version(
    session: AsyncSession,
    ds: Dataset,
    table: pa.Table,
    *,
    step: dict[str, Any],
    summary: str,
    actor_id: uuid.UUID,
    parent: DatasetVersion | None,
) -> DatasetVersion:
    """Append an immutable version. The dataset row is locked so two writers queue."""
    await session.refresh(ds, with_for_update=True)
    number = ds.latest_version + 1
    key, digest, _ = await storage.put_table(ds.id, number, table)
    lineage = list(parent.lineage) if parent else []
    lineage.append({**step, "at": utcnow().isoformat(), "by": str(actor_id)})
    v = DatasetVersion(
        dataset_id=ds.id,
        number=number,
        parent_version_id=parent.id if parent else None,
        storage_key=key,
        sha256=digest,
        row_count=table.num_rows,
        columns=frame.schema_of(table),
        lineage=lineage,
        summary=summary[:300],
        created_by=actor_id,
    )
    session.add(v)
    ds.latest_version = number
    await session.flush()
    return v


async def create_dataset(
    session: AsyncSession,
    principal: Principal,
    *,
    name: str,
    table: pa.Table,
    description: str | None = None,
    tags: list[str] | None = None,
    step: dict[str, Any],
    summary: str,
    rules: list[dict[str, Any]] | None = None,
    collection_key: str | None = None,
    parent: DatasetVersion | None = None,
) -> tuple[Dataset, DatasetVersion]:
    ds = Dataset(
        name=name.strip(),
        description=description,
        tags=tags or [],
        owner_id=principal.user_id,
        validation_rules=rules or [],
        collection_key=collection_key,
    )
    session.add(ds)
    try:
        await session.flush()
    except IntegrityError as exc:
        await session.rollback()
        raise Conflict("A dataset with that collection key already exists.") from exc
    v = await write_version(
        session, ds, table, step=step, summary=summary, actor_id=principal.user_id, parent=parent
    )
    await record(
        session,
        audit_context(principal),
        action="create",
        entity_type="dataset",
        entity_id=ds.id,
        after={"name": ds.name, "rows": v.row_count, "columns": len(v.columns)},
        only_changes=False,
    )
    return ds, v


# ------------------------------------------------------------------ import


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValidationFailed("The file text could not be decoded.")


def read_tabular(data: bytes, kind: str) -> tuple[list[str], list[list[Any]]] | pa.Table:
    """Header and text rows for csv, xlsx and json; an Arrow table for parquet."""
    if kind == "parquet":
        return pq.read_table(io.BytesIO(data))
    if kind == "csv":
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif kind == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        rows = [
            ["" if c is None else c for c in r]
            for r in wb.worksheets[0].iter_rows(values_only=True)
        ]
    elif kind == "json":
        import json

        doc = json.loads(_decode(data))
        if not isinstance(doc, list) or not doc or not all(isinstance(x, dict) for x in doc):
            raise ValidationFailed("JSON must be a list of objects, one per row.")
        header = list(dict.fromkeys(k for r in doc for k in r))
        return header, [[r.get(h) for h in header] for r in doc]
    else:
        raise ValidationFailed(f"Files of type {kind} cannot be imported as data.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValidationFailed("The file has no rows.")
    header = [str(h).strip() or f"column_{i + 1}" for i, h in enumerate(rows[0])]
    if len(set(header)) != len(header):
        raise ValidationFailed("Two columns have the same name. Rename one in the file.")
    if len(header) > frame.MAX_COLUMNS:
        raise ValidationFailed(f"A dataset can have at most {frame.MAX_COLUMNS} columns.")
    return header, rows[1:]


async def preview_import(
    session: AsyncSession, principal: Principal, file_id: uuid.UUID
) -> dict[str, Any]:
    principal.require(P.DATASET_WRITE)
    ref, data = await read_file(session, principal, file_id)
    if ref.purpose not in ("dataset_import", "library"):
        raise ValidationFailed("Upload the file with the purpose dataset_import first.")
    parsed = read_tabular(data, ref.kind)
    if isinstance(parsed, pa.Table):
        return {
            "file": ref.original_name,
            "rows": parsed.num_rows,
            "schema": [{**c, "confidence": 1.0} for c in frame.schema_of(parsed)],
            "sample": frame.rows(parsed, 20),
        }
    header, rows = parsed
    schema = infer.infer_schema(header, rows)
    return {
        "file": ref.original_name,
        "rows": len(rows),
        "schema": schema,
        "sample": [
            dict(zip(header, [None if c == "" else str(c) for c in r], strict=False))
            for r in rows[:20]
        ],
    }


def _build(
    parsed: Any, schema_override: list[dict[str, str]] | None
) -> tuple[pa.Table, list[dict[str, Any]]]:
    if isinstance(parsed, pa.Table):
        return parsed, []
    header, rows = parsed
    schema = infer.infer_schema(header, rows)
    if schema_override:
        by = {c["name"]: c["type"] for c in schema_override}
        bad = [n for n in by if n not in header]
        if bad:
            raise ValidationFailed(f"The schema names unknown columns: {', '.join(bad)}.")
        if any(t not in frame.TYPES for t in by.values()):
            raise ValidationFailed("Unknown column type in the schema.")
        schema = [{**c, "type": by.get(c["name"], c["type"])} for c in schema]
    return infer.build_table(header, rows, schema)


async def create_from_import(
    session: AsyncSession,
    principal: Principal,
    *,
    file_id: uuid.UUID,
    name: str,
    description: str | None,
    tags: list[str],
    schema: list[dict[str, str]] | None,
    rules: list[dict[str, Any]],
) -> tuple[Dataset, DatasetVersion, int]:
    principal.require(P.DATASET_WRITE)
    ref, data = await read_file(session, principal, file_id)
    if ref.purpose not in ("dataset_import", "library"):
        raise ValidationFailed("Upload the file with the purpose dataset_import first.")
    table, bad = _build(read_tabular(data, ref.kind), schema)
    parsed_rules = _RULES.validate_python(rules)
    try:
        table, bad_rules = apply_rules(table, parsed_rules)
    except ValueError as exc:
        raise ValidationFailed(str(exc)) from exc
    if table.num_rows == 0:
        raise ValidationFailed(
            "No rows are left after checking types and rules. See the quarantine preview.",
            code="all_rows_bad",
        )
    ds, v = await create_dataset(
        session,
        principal,
        name=name,
        table=table,
        description=description,
        tags=tags,
        step={"op": "import", "file": ref.original_name, "sha256": ref.sha256},
        summary=f"Imported {ref.original_name}",
        rules=rules,
    )
    quarantined = 0
    for b in [*bad, *bad_rules]:
        session.add(
            QuarantineRow(
                dataset_id=ds.id,
                version_id=v.id,
                row_index=b["row_index"],
                errors=b["errors"],
                raw=b["raw"],
            )
        )
        quarantined += 1
    await session.commit()
    await session.refresh(ds)
    await session.refresh(v)
    return ds, v, quarantined


async def create_from_rows(
    session: AsyncSession,
    principal: Principal,
    *,
    name: str,
    description: str | None,
    tags: list[str],
    columns: list[dict[str, str]],
    rows: list[list[Any]],
) -> tuple[Dataset, DatasetVersion, int]:
    principal.require(P.DATASET_WRITE)
    header = [c["name"] for c in columns]
    if len(set(header)) != len(header):
        raise ValidationFailed("Column names must be different.")
    table, bad = infer.build_table(header, rows, columns)
    if table.num_rows == 0:
        raise ValidationFailed("No valid rows were given.", code="all_rows_bad")
    ds, v = await create_dataset(
        session,
        principal,
        name=name,
        table=table,
        description=description,
        tags=tags,
        step={"op": "manual_entry"},
        summary="Entered by hand",
    )
    for b in bad:
        session.add(
            QuarantineRow(
                dataset_id=ds.id,
                version_id=v.id,
                row_index=b["row_index"],
                errors=b["errors"],
                raw=b["raw"],
            )
        )
    await session.commit()
    return ds, v, len(bad)


async def append_rows(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    rows: list[list[Any]],
    version: int,
) -> tuple[DatasetVersion, int]:
    ds = await get_dataset(session, principal, dataset_id, edit=True)
    if ds.version != version:
        raise StaleVersion()
    latest = await get_version(session, ds)
    current = await load_table(latest)
    header = current.column_names
    schema = [{"name": c["name"], "type": c["type"]} for c in latest.columns]
    added, bad = infer.build_table(header, rows, schema)
    added, bad_rules = apply_rules(added, _RULES.validate_python(ds.validation_rules))
    merged = pa.concat_tables([current, added.cast(current.schema)])
    v = await write_version(
        session,
        ds,
        merged,
        step={"op": "append_rows", "rows": added.num_rows},
        summary=f"Added {added.num_rows} rows",
        actor_id=principal.user_id,
        parent=latest,
    )
    for b in [*bad, *bad_rules]:
        session.add(
            QuarantineRow(
                dataset_id=ds.id,
                version_id=v.id,
                row_index=b["row_index"],
                errors=b["errors"],
                raw=b["raw"],
            )
        )
    await record(
        session,
        audit_context(principal),
        action="append_rows",
        entity_type="dataset",
        entity_id=ds.id,
        after={
            "rows_added": added.num_rows,
            "quarantined": len(bad) + len(bad_rules),
            "version": v.number,
        },
        only_changes=False,
    )
    await session.commit()
    await session.refresh(v)
    return v, len(bad) + len(bad_rules)


# ------------------------------------------------------------------ pipeline and analysis


async def synonyms_map(session: AsyncSession) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    for s in await session.scalars(select(Synonym)):
        out.setdefault(s.domain, {})[s.alias.strip().lower()] = s.canonical
    return out


async def run_steps(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    *,
    version: int | None,
    raw_steps: list[dict[str, Any]],
    save_as: str | None,
    dry_run: bool,
) -> tuple[pa.Table, DatasetVersion | None]:
    ds = await get_dataset(session, principal, dataset_id, edit=not dry_run)
    src = await get_version(session, ds, version)
    try:
        steps = _STEPS.validate_python(raw_steps)
    except ValueError as exc:
        raise ValidationFailed(
            f"The steps are not valid: {str(exc).splitlines()[0]}", code="steps_invalid"
        ) from exc
    base = await load_table(src)
    others: dict[uuid.UUID, pa.Table] = {}
    for s in steps:
        if isinstance(s, Join):
            ov = await session.get(DatasetVersion, s.other_version_id)
            if ov is None:
                raise ValidationFailed("A joined dataset version does not exist.")
            await get_dataset(session, principal, ov.dataset_id)  # the caller must see it too
            others[ov.id] = await load_table(ov)
    try:
        result = pipeline.run_pipeline(
            base, steps, others=others, synonyms=await synonyms_map(session)
        )
    except pipeline.StepError as exc:
        raise ValidationFailed(str(exc), code="step_failed", extra={"step": exc.index + 1}) from exc
    if dry_run:
        return result, None
    lineage_step = {
        "op": "pipeline",
        "steps": [s.model_dump(mode="json") for s in steps],
        "source_version": src.number,
    }
    if save_as:
        new_ds, v = await create_dataset(
            session,
            principal,
            name=save_as,
            table=result,
            step=lineage_step,
            summary=f"{len(steps)} step(s) applied to {ds.name} v{src.number}",
            parent=src,
        )
        target = new_ds
    else:
        v = await write_version(
            session,
            ds,
            result,
            step=lineage_step,
            summary=f"{len(steps)} step(s) applied",
            actor_id=principal.user_id,
            parent=src,
        )
        target = ds
    await record(
        session,
        audit_context(principal),
        action="run_pipeline",
        entity_type="dataset",
        entity_id=target.id,
        after={
            "steps": len(steps),
            "rows": v.row_count,
            "version": v.number,
            "from": f"{ds.id}:v{src.number}",
        },
        only_changes=False,
    )
    await session.commit()
    if v is not None:
        await session.refresh(v)
    return result, v


async def analysis(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    *,
    outlier_z: float,
    iqr_k: float,
) -> dict[str, Any]:
    ds = await get_dataset(session, principal, dataset_id)
    v = await get_version(session, ds, version)
    key = eda.config_hash({"z": outlier_z, "k": iqr_k})
    if key in v.analysis_cache:
        return dict(v.analysis_cache[key])
    table = await load_table(v)
    result = {
        "rows": v.row_count,
        "columns": eda.column_stats(table, outlier_z=outlier_z, iqr_k=iqr_k),
        "correlations": eda.correlations(table),
    }
    v.analysis_cache = {**v.analysis_cache, key: result}  # cached per version and settings
    await session.commit()
    return result


async def chart(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    **spec: Any,
) -> dict[str, Any]:
    ds = await get_dataset(session, principal, dataset_id)
    table = await load_table(await get_version(session, ds, version))
    try:
        return charts.build(table, **spec)
    except ValueError as exc:
        raise ValidationFailed(str(exc)) from exc


async def preview_rows(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    offset: int,
    limit: int,
) -> dict[str, Any]:
    ds = await get_dataset(session, principal, dataset_id)
    v = await get_version(session, ds, version)
    table = await load_table(v)
    limit = min(limit, frame.PREVIEW_CAP)
    return {
        "version": v.number,
        "total": v.row_count,
        "columns": v.columns,
        "rows": frame.rows(table, limit, offset),
    }


async def diff_versions(
    session: AsyncSession, principal: Principal, dataset_id: uuid.UUID, a: int, b: int
) -> dict[str, Any]:
    ds = await get_dataset(session, principal, dataset_id)
    va, vb = await get_version(session, ds, a), await get_version(session, ds, b)
    ca, cb = {c["name"]: c["type"] for c in va.columns}, {c["name"]: c["type"] for c in vb.columns}
    out: dict[str, Any] = {
        "from": a,
        "to": b,
        "rows": {"from": va.row_count, "to": vb.row_count, "change": vb.row_count - va.row_count},
        "columns_added": [c for c in cb if c not in ca],
        "columns_removed": [c for c in ca if c not in cb],
        "type_changes": [
            {"column": c, "from": ca[c], "to": cb[c]} for c in ca if c in cb and ca[c] != cb[c]
        ],
        "numeric": [],
    }
    shared = [c for c in ca if c in cb and ca[c] in eda.NUMERIC and cb[c] in eda.NUMERIC]
    if shared:
        sa = {s["name"]: s for s in eda.column_stats((await load_table(va)).select(shared))}
        sb = {s["name"]: s for s in eda.column_stats((await load_table(vb)).select(shared))}
        for c in shared:
            out["numeric"].append(
                {
                    "column": c,
                    **{
                        k: {"from": sa[c].get(k), "to": sb[c].get(k)}
                        for k in ("mean", "min", "max")
                    },
                }
            )
    return out


# ------------------------------------------------------------------ management


async def update_dataset(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    *,
    version: int,
    name: str | None,
    description: str | None,
    tags: list[str] | None,
    rules: list[dict[str, Any]] | None,
) -> Dataset:
    ds = await get_dataset(session, principal, dataset_id, edit=True)
    if ds.version != version:
        raise StaleVersion()
    before = {"name": ds.name, "tags": ds.tags}
    if name is not None:
        ds.name = name.strip()
    if description is not None:
        ds.description = description
    if tags is not None:
        ds.tags = tags
    if rules is not None:
        try:
            _RULES.validate_python(rules)
        except ValueError as exc:
            raise ValidationFailed("A validation rule is not valid.") from exc
        ds.validation_rules = rules
    await record(
        session,
        audit_context(principal),
        action="update",
        entity_type="dataset",
        entity_id=ds.id,
        before=before,
        after={"name": ds.name, "tags": ds.tags},
    )
    await session.commit()
    await session.refresh(ds)
    return ds


async def duplicate(
    session: AsyncSession, principal: Principal, dataset_id: uuid.UUID, name: str
) -> Dataset:
    src = await get_dataset(session, principal, dataset_id)
    principal.require(P.DATASET_WRITE)
    table = await load_table(await get_version(session, src))
    ds, _ = await create_dataset(
        session,
        principal,
        name=name,
        table=table,
        description=src.description,
        tags=list(src.tags),
        step={"op": "duplicate", "of": str(src.id)},
        summary=f"Copy of {src.name}",
    )
    await session.commit()
    await session.refresh(ds)
    return ds


async def soft_delete(session: AsyncSession, principal: Principal, dataset_id: uuid.UUID) -> None:
    ds = await get_dataset(session, principal, dataset_id)
    _owner_or_admin(principal, ds)
    if ds.collection_key:
        raise Conflict(
            "Built-in library collections cannot be deleted.", code="collection_protected"
        )
    ds.deleted_at = utcnow()
    await record(
        session,
        audit_context(principal),
        action="delete",
        entity_type="dataset",
        entity_id=ds.id,
        before={"name": ds.name},
        only_changes=False,
    )
    await session.commit()


async def restore(session: AsyncSession, principal: Principal, dataset_id: uuid.UUID) -> Dataset:
    ds = await get_dataset(session, principal, dataset_id, deleted=True)
    _owner_or_admin(principal, ds)
    ds.deleted_at = None
    await record(
        session,
        audit_context(principal),
        action="restore",
        entity_type="dataset",
        entity_id=ds.id,
        after={"name": ds.name},
        only_changes=False,
    )
    await session.commit()
    await session.refresh(ds)
    return ds


async def share(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    *,
    kind: str,
    target: str,
    access: str,
) -> DatasetShare:
    ds = await get_dataset(session, principal, dataset_id)
    _owner_or_admin(principal, ds)
    if kind not in ("user", "role") or access not in ("view", "edit"):
        raise ValidationFailed("Share with a user or a role, as view or edit.")
    row = await session.get(DatasetShare, (ds.id, kind, target))
    if row is None:
        row = DatasetShare(
            dataset_id=ds.id,
            kind=kind,
            principal=target,
            access=access,
            granted_by=principal.user_id,
        )
        session.add(row)
    else:
        row.access = access
    await record(
        session,
        audit_context(principal),
        action="share",
        entity_type="dataset",
        entity_id=ds.id,
        after={"kind": kind, "target": target, "access": access},
        only_changes=False,
    )
    await session.commit()
    return row


async def unshare(
    session: AsyncSession, principal: Principal, dataset_id: uuid.UUID, *, kind: str, target: str
) -> None:
    ds = await get_dataset(session, principal, dataset_id)
    _owner_or_admin(principal, ds)
    row = await session.get(DatasetShare, (ds.id, kind, target))
    if row is None:
        raise NotFound("That share does not exist.")
    await session.delete(row)
    await record(
        session,
        audit_context(principal),
        action="unshare",
        entity_type="dataset",
        entity_id=ds.id,
        before={"kind": kind, "target": target},
        only_changes=False,
    )
    await session.commit()


async def list_shares(
    session: AsyncSession, principal: Principal, dataset_id: uuid.UUID
) -> list[DatasetShare]:
    ds = await get_dataset(session, principal, dataset_id)
    _owner_or_admin(principal, ds)
    return list(await session.scalars(select(DatasetShare).where(DatasetShare.dataset_id == ds.id)))


async def freeze_for_training(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    *,
    source: str,
    known_issues: str | None,
) -> DatasetVersion:
    ds = await get_dataset(session, principal, dataset_id, edit=True)
    v = await get_version(session, ds, version)
    if v.frozen_for_training:
        raise Conflict("This version is already a frozen training snapshot.")
    table = await load_table(v)
    date_cols = [c["name"] for c in v.columns if c["type"] in ("date", "timestamp")]
    span = None
    if date_cols:
        st = eda.column_stats(table.select(date_cols[:1]))[0]
        span = {"column": date_cols[0], "from": st.get("min"), "to": st.get("max")}
    # Label balance and cleaning provenance, when the dataset carries them (the BOQ collection
    # does, ADR 0014). A model card reads these to say what the model saw.
    distributions: dict[str, dict[str, int]] = {}
    for col in ("gap_type", "line_role", "cleaning_version", "doc_kind"):
        if col in table.column_names:
            counts: dict[str, int] = {}
            for val in table.column(col).to_pylist():
                k = "missing" if val is None else str(val)
                counts[k] = counts.get(k, 0) + 1
            distributions[col] = dict(sorted(counts.items()))
    v.frozen_for_training = True
    v.data_card = {
        "dataset": ds.name,
        "version": v.number,
        "source": source,
        "row_count": v.row_count,
        "columns": v.columns,
        "distributions": distributions,
        "date_range": span,
        "known_issues": known_issues,
        "sha256": v.sha256,
        "frozen_at": utcnow().isoformat(),
        "frozen_by": str(principal.user_id),
    }
    await record(
        session,
        audit_context(principal),
        action="freeze_training",
        entity_type="dataset",
        entity_id=ds.id,
        after={"version": v.number, "sha256": v.sha256},
        only_changes=False,
    )
    await session.commit()
    return v


# ------------------------------------------------------------------ quarantine


async def list_quarantine(
    session: AsyncSession, principal: Principal, dataset_id: uuid.UUID, status: str = "open"
) -> list[QuarantineRow]:
    ds = await get_dataset(session, principal, dataset_id)
    rows = await session.scalars(
        select(QuarantineRow)
        .where(QuarantineRow.dataset_id == ds.id, QuarantineRow.status == status)
        .order_by(QuarantineRow.created_at, QuarantineRow.row_index)
        .limit(500)
    )
    return list(rows)


async def resolve_quarantine(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    row_id: uuid.UUID,
    *,
    action: str,
    values: dict[str, Any] | None,
) -> QuarantineRow:
    # Anyone who may write datasets can curate the built-in library collections; other datasets
    # need edit access as usual.
    ds = await get_dataset(session, principal, dataset_id)
    principal.require(P.DATASET_WRITE)
    if ds.collection_key is None:
        ds = await get_dataset(session, principal, dataset_id, edit=True)
    q = await session.get(QuarantineRow, row_id)
    if q is None or q.dataset_id != ds.id or q.status != "open":
        raise NotFound("That held row was not found or is already resolved.")
    if action == "discard":
        q.status = "discarded"
    elif action == "fix":
        merged = {**q.raw, **(values or {})}
        if "confidence" in merged:
            merged["confidence"] = "1"  # a person has now checked this row
        latest = await get_version(session, ds)
        cur = await load_table(latest)
        schema = [{"name": c["name"], "type": c["type"]} for c in latest.columns]
        fixed, bad = infer.build_table(
            cur.column_names, [[merged.get(h) for h in cur.column_names]], schema
        )
        fixed, bad_rules = apply_rules(fixed, _RULES.validate_python(ds.validation_rules))
        if bad or bad_rules or fixed.num_rows != 1:
            raise ValidationFailed(
                "The corrected row still has problems: "
                + "; ".join((bad or bad_rules)[0]["errors"])
            )
        v = await write_version(
            session,
            ds,
            pa.concat_tables([cur, fixed.cast(cur.schema)]),
            step={"op": "fix_held_row", "row": q.row_index},
            summary="Corrected a held row",
            actor_id=principal.user_id,
            parent=latest,
        )
        q.status, q.version_id = "fixed", v.id
    else:
        raise ValidationFailed("Action must be fix or discard.")
    q.resolved_by = principal.user_id
    if q.library_file_id is not None:
        # The last held line of a BOQ file is settled: the file no longer needs a person.
        await session.flush()
        still_open = await session.scalar(
            select(func.count())
            .select_from(QuarantineRow)
            .where(
                QuarantineRow.library_file_id == q.library_file_id, QuarantineRow.status == "open"
            )
        )
        lf = await session.get(LibraryFile, q.library_file_id)
        if lf is not None and not still_open and lf.kind == "boq" and lf.status == "needs_review":
            lf.status = "processed"
            lf.message = f"{lf.rows_added} lines added; held lines settled by a person."
    await record(
        session,
        audit_context(principal),
        action=f"quarantine_{action}",
        entity_type="dataset",
        entity_id=ds.id,
        after={"row": q.row_index},
        only_changes=False,
    )
    await session.commit()
    return q


# ------------------------------------------------------------------ export


async def export(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    fmt: str,
) -> tuple[bytes, str, str]:
    ds = await get_dataset(session, principal, dataset_id)
    v = await get_version(session, ds, version)
    table = await load_table(v)
    base = f"{ds.name.replace(' ', '_')[:60]}_v{v.number}"
    if fmt == "parquet":
        data, mime = frame.to_parquet(table), "application/vnd.apache.parquet"
    elif fmt == "csv":
        import pyarrow.csv as pacsv

        sink = io.BytesIO()
        pacsv.write_csv(table, sink)
        data, mime = sink.getvalue(), "text/csv"
    elif fmt == "xlsx":
        if table.num_rows > EXPORT_ROWS_XLSX:
            raise ValidationFailed(
                f"Excel export is limited to {EXPORT_ROWS_XLSX:,} rows. Use CSV or Parquet."
            )
        import openpyxl

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet("data")
        ws.append(table.column_names)
        for r in table.to_pylist():
            ws.append([_cell(x) for x in r.values()])
        sink = io.BytesIO()
        wb.save(sink)
        data, mime = (
            sink.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        raise ValidationFailed("Export as csv, xlsx or parquet.")
    await record(
        session,
        audit_context(principal),
        action="export",
        entity_type="dataset",
        entity_id=ds.id,
        after={"format": fmt, "rows": v.row_count, "version": v.number},
        only_changes=False,
    )
    await session.commit()
    return data, mime, f"{base}.{fmt}"


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, bool | int | float | str):
        return v
    if hasattr(v, "tzinfo") and getattr(v, "tzinfo", None) is not None:
        return v.replace(tzinfo=None)
    return v if hasattr(v, "isoformat") else str(v)


# ------------------------------------------------------------------ synonyms and promotion


async def add_synonym(
    session: AsyncSession, principal: Principal, *, domain: str, canonical: str, alias: str
) -> Synonym:
    principal.require(P.CATALOGUE_WRITE)
    s = Synonym(domain=domain, canonical=canonical.strip(), alias=alias.strip())
    session.add(s)
    try:
        await session.flush()
    except IntegrityError as exc:
        await session.rollback()
        raise Conflict("That alias already exists for this domain.") from exc
    await record(
        session,
        audit_context(principal),
        action="add_synonym",
        entity_type="synonym",
        entity_id=s.id,
        after={"domain": domain, "canonical": canonical, "alias": alias},
        only_changes=False,
    )
    await session.commit()
    return s


async def list_synonyms(session: AsyncSession, principal: Principal) -> list[Synonym]:
    principal.require(P.DATASET_READ)
    return list(
        await session.scalars(
            select(Synonym).order_by(Synonym.domain, Synonym.canonical, Synonym.alias)
        )
    )


async def submit_promotions(
    session: AsyncSession,
    principal: Principal,
    dataset_id: uuid.UUID,
    version: int | None,
    *,
    mapping: dict[str, str],
    constants: dict[str, Any],
    row_numbers: list[int] | None,
) -> list[PromotionItem]:
    ds = await get_dataset(session, principal, dataset_id, edit=True)
    v = await get_version(session, ds, version)
    table = await load_table(v)
    missing = [c for c in mapping.values() if c not in table.column_names]
    if missing:
        raise ValidationFailed(f"Unknown columns in the mapping: {', '.join(missing)}.")
    rows = table.to_pylist()
    picked = [
        (i + 1, r) for i, r in enumerate(rows) if row_numbers is None or (i + 1) in set(row_numbers)
    ]
    if len(picked) > MAX_PROMOTE:
        raise ValidationFailed(f"Promote at most {MAX_PROMOTE} rows at a time.")
    out: list[PromotionItem] = []
    errors: list[str] = []
    for num, r in picked:
        draft = {
            **constants,
            **{
                field: frame._jsonable(r[col])
                for field, col in mapping.items()
                if r[col] is not None
            },
        }
        problems = catalogue.check_item_draft(draft)
        if problems:
            errors.append(f"Row {num}: {problems[0]}")
            continue
        item = PromotionItem(
            dataset_id=ds.id,
            version_id=v.id,
            row={k: frame._jsonable(x) for k, x in r.items()},
            proposed=draft,
            submitted_by=principal.user_id,
        )
        session.add(item)
        out.append(item)
    if errors and not out:
        raise ValidationFailed(
            "No row is ready for the catalogue. " + errors[0], code="nothing_promotable"
        )
    await record(
        session,
        audit_context(principal),
        action="submit_promotions",
        entity_type="dataset",
        entity_id=ds.id,
        after={"submitted": len(out), "skipped": len(errors)},
        only_changes=False,
    )
    await session.commit()
    return out


async def list_promotions(
    session: AsyncSession, principal: Principal, status: str = "pending"
) -> list[PromotionItem]:
    principal.require(P.CATALOGUE_WRITE)
    rows = await session.scalars(
        select(PromotionItem)
        .where(PromotionItem.status == status)
        .order_by(PromotionItem.created_at)
        .limit(500)
    )
    return list(rows)


async def decide_promotion(
    session: AsyncSession,
    principal: Principal,
    promotion_id: uuid.UUID,
    *,
    approve: bool,
    note: str | None,
) -> PromotionItem:
    principal.require(P.CATALOGUE_WRITE)
    p = await session.get(PromotionItem, promotion_id)
    if p is None or p.status != "pending":
        raise NotFound("That proposal was not found or is already decided.")
    ensure_different_people(p.submitted_by, principal.user_id, "catalogue proposal")
    if approve:
        try:
            item = await catalogue.create_item_from_draft(session, principal, p.proposed)
        except (ValueError, Conflict) as exc:
            raise ValidationFailed(f"The catalogue refused this item: {exc}") from exc
        p.catalogue_item_id = item.id
        p.status = "approved"
    else:
        if not note or len(note.strip()) < 3:
            raise ValidationFailed("Say why the proposal is rejected.")
        p.status = "rejected"
    p.reviewed_by, p.reviewed_at, p.note = principal.user_id, utcnow(), note
    await record(
        session,
        audit_context(principal),
        action="approve_promotion" if approve else "reject_promotion",
        entity_type="dataset_promotion",
        entity_id=p.id,
        after={"item": str(p.catalogue_item_id) if p.catalogue_item_id else None, "note": note},
        only_changes=False,
    )
    await session.commit()
    return p
