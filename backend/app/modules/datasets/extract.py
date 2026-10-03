"""Plain text out of PDF, Word, Excel and JSON files, page by page.

This is the only place in the corpus pipeline that opens the heavy original. Everything after it
works on the text and on the parser's structured result."""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "unknown": "application/octet-stream",
}


@dataclass(frozen=True)
class Extracted:
    file_kind: str  # pdf | docx | xlsx | json | unknown
    pages: list[str]  # one entry per PDF page, Excel sheet, or the whole Word document

    @property
    def chars(self) -> int:
        return sum(len(p) for p in self.pages)


def sniff(data: bytes) -> str:
    """The file kind from its bytes, never from its name."""
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return "unknown"
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        return "unknown"
    stripped = data.lstrip()[:1]
    if stripped in (b"{", b"["):
        try:
            json.loads(data)
        except ValueError:
            return "unknown"
        return "json"
    return "unknown"


def _pdf(data: bytes) -> list[str]:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [(p.extract_text() or "").replace("\r", "") for p in pdf.pages]


def _docx(data: bytes) -> list[str]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            cells: list[str] = []
            for c in row.cells:
                text = c.text.strip()
                if text and (not cells or cells[-1] != text):  # merged cells repeat
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return ["\n".join(parts)]


def _xlsx(data: bytes) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                rows.append(" | ".join(vals))
        pages.append("\n".join(rows))
    wb.close()
    return pages


def extract(data: bytes) -> Extracted:
    kind = sniff(data)
    if kind == "pdf":
        return Extracted(kind, _pdf(data))
    if kind == "docx":
        return Extracted(kind, _docx(data))
    if kind == "xlsx":
        return Extracted(kind, _xlsx(data))
    if kind == "json":
        return Extracted(kind, [json.dumps(json.loads(data), ensure_ascii=False, indent=1)])
    return Extracted(kind, [])
