"""Historical BOQ importers: the ITCraft priced quotation and summary BOQ, from PDF or XLSX.

The output is a flat list of lines plus document facts, each line carrying a confidence score.
Low-confidence lines are never dropped and never trusted silently: the library sends them to a
human review queue. Nothing here touches the database.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.core.money import parse_inr

GROUPS = ("high priority", "to consider")
_AMOUNT = r"₹\s*([\d,]+(?:\.\d{1,2})?)"
_PRICED = re.compile(
    rf"^(?P<ref>\d+[A-Z]?)\s+(?P<text>.+?)\s+(?P<qty>\d+)\s+{_AMOUNT}\s+{_AMOUNT}\s*$"
)
_SUMMARY = re.compile(r"^(?P<ref>\d+[A-Z]?)\s+(?P<text>.+?)\s+(?P<qty>\d+)\s*$")
_QUOTE_REF = re.compile(r"Quote\s*Ref\s*No\s*:\s*([A-Za-z0-9/\-]+)")
_DATE = re.compile(r"Date\s*:\s*(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+),?\s+(\d{4})")
_VALID = re.compile(r"valid\s+for\s+(?:a\s+)?period\s+of\s+(\d+)\s+days?", re.I)
_GST = re.compile(r"GST\s*@\s*(\d+(?:\.\d+)?)\s*%", re.I)
_GARBLED = re.compile(r"[A-Z][A-Z][a-z]")  # interleaved page-header text, e.g. "FHoigrhti"
_OPTION = re.compile(r"^(\d+)([A-Z])$")

MONTHS = {
    m: i
    for i, m in enumerate(
        [
            "january",
            "february",
            "march",
            "april",
            "may",
            "june",
            "july",
            "august",
            "september",
            "october",
            "november",
            "december",
        ],
        start=1,
    )
}


@dataclass
class BoqLine:
    line_ref: str
    option: str
    priority_group: str
    section: str
    component: str
    description: str
    inclusions: list[str]
    qty: int
    unit_price: Decimal | None
    amount: Decimal | None
    confidence: float
    issues: list[str] = field(default_factory=list)


@dataclass
class BoqDocument:
    kind: str  # "quotation" or "summary"
    quote_ref: str | None
    quote_date: date | None
    customer: str | None
    validity_days: int | None
    gst_rate: Decimal | None
    lines: list[BoqLine]
    warnings: list[str] = field(default_factory=list)


class BoqParseError(Exception):
    pass


def _to_decimal(s: str) -> Decimal | None:
    try:
        return parse_inr(s)
    except (ValueError, InvalidOperation):
        return None


def _parse_date(text: str) -> date | None:
    m = _DATE.search(text)
    if not m:
        return None
    month = MONTHS.get(m.group(2).lower())
    if not month:
        return None
    try:
        return date(int(m.group(3)), month, int(m.group(1)))
    except ValueError:
        return None


def _customer(lines: list[str]) -> str | None:
    for i, line in enumerate(lines):
        if line.strip() == "To,":
            for nxt in lines[i + 1 : i + 3]:
                m = re.match(r"^(?:M/S\.?|Messrs\.?)\s+(.*?)(?:\s+Date\s*:.*)?$", nxt.strip(), re.I)
                if m:
                    return m.group(1).strip()
    return None


def _split_lines(text: str) -> list[str]:
    return [ln.rstrip() for ln in text.replace("\r", "").split("\n") if ln.strip()]


_SECTION_WORDS = re.compile(
    r"(options?|alternatives?|services?|licen[cs]es?|hardware|software)\s*$", re.I
)


def _looks_like_section(ln: str, cur: BoqLine) -> bool:
    """A heading between items, such as 'Firewall Options', versus a wrapped description line.
    Headings are short, start with a capital, have no sentence punctuation and no digits."""
    if len(ln) > 70 or re.search(r"[.:;/]|\d", ln):
        return False
    if not ln[:1].isupper():
        return False
    return bool(_SECTION_WORDS.search(ln)) or ("," in ln and "&" in ln)


def parse_lines(lines: list[str]) -> BoqDocument:
    """Parse already-extracted text lines (the PDF and the tests both come through here)."""
    joined = "\n".join(lines)
    quote_ref = m.group(1) if (m := _QUOTE_REF.search(joined)) else None
    quote_date = _parse_date(joined)
    customer = _customer(lines)
    validity = int(m.group(1)) if (m := _VALID.search(joined)) else None
    gst = Decimal(m.group(1)) if (m := _GST.search(joined)) else None

    start = next(
        (i for i, ln in enumerate(lines) if re.match(r"^Sr\.?\s*N", ln.strip(), re.I)), None
    )
    if start is None:
        raise BoqParseError("No BOQ table header (Sr N, Components, Qty) was found.")
    priced = "Price" in lines[start]
    stop = next(
        (i for i, ln in enumerate(lines) if ln.strip().lower().startswith("terms")), len(lines)
    )

    out: list[BoqLine] = []
    warnings: list[str] = []
    group, section = "", ""
    cur: BoqLine | None = None
    rx = _PRICED if priced else _SUMMARY

    for raw in lines[start + 1 : stop]:
        ln = raw.strip()
        if re.match(r"^Sr\.?\s*N", ln, re.I):  # header repeated on the next page
            continue
        m = rx.match(ln)
        if m:
            ref = m.group("ref")
            text = m.group("text").strip()
            qty = int(m.group("qty"))
            price = amount = None
            issues: list[str] = []
            conf = 1.0
            if priced:
                price, amount = _to_decimal(m.group(4)), _to_decimal(m.group(5))
                if price is None or amount is None:
                    issues.append("Price or amount could not be read.")
                    conf = 0.4
                elif price * qty != amount:
                    issues.append(f"Amount {amount} is not quantity {qty} times price {price}.")
                    conf = min(conf, 0.5)
            if _GARBLED.search(text.split(" ")[0] if text else ""):
                issues.append(
                    "The component text looks garbled (overlapping print on a page break)."
                )
                conf = min(conf, 0.5)
            om = _OPTION.match(ref)
            cur = BoqLine(
                line_ref=ref,
                option=om.group(2) if om else "",
                priority_group=group,
                section=section,
                component=text,
                description="",
                inclusions=[],
                qty=qty,
                unit_price=price,
                amount=amount,
                confidence=conf,
                issues=issues,
            )
            out.append(cur)
            continue
        if ln.lower() in GROUPS:
            group, section, cur = ln.title(), "", None
            continue
        if ln.startswith(("•", "*", "-")) and cur is not None:
            cur.inclusions.append(ln.lstrip("•*- ").strip())
            continue
        if cur is not None and not cur.inclusions and not _looks_like_section(ln, cur):
            # continuation of the title or a one-line description under the item
            cur.description = f"{cur.description} {ln}".strip()
            continue
        section, cur = ln, None

    if not out:
        raise BoqParseError("The table was found but no lines could be read from it.")
    return BoqDocument(
        kind="quotation" if priced else "summary",
        quote_ref=quote_ref,
        quote_date=quote_date,
        customer=customer,
        validity_days=validity,
        gst_rate=gst,
        lines=out,
        warnings=warnings,
    )


def parse_pdf(data: bytes) -> BoqDocument:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise BoqParseError("pdfplumber is not installed") from exc
    lines: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                lines.extend(_split_lines(page.extract_text() or ""))
    except Exception as exc:
        raise BoqParseError(f"The PDF could not be read: {exc}") from exc
    return parse_lines(lines)


def detect_pdf(data: bytes) -> float:
    """Confidence that a PDF is a BOQ or quotation. Never raises."""
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            first = (pdf.pages[0].extract_text() or "") if pdf.pages else ""
    except Exception:
        return 0.0
    low = first.lower()
    score = 0.0
    if "boq" in low or "quotation" in low:
        score += 0.4
    if re.search(r"sr\.?\s*n", low) and "components" in low and "qty" in low:
        score += 0.5
    return min(score, 1.0)


def parse_xlsx(data: bytes) -> BoqDocument:
    """A spreadsheet BOQ: find the row with Components and Qty, read rows below it."""
    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise BoqParseError(f"The spreadsheet could not be read: {exc}") from exc
    ws = wb.worksheets[0]
    rows = [
        [("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)
    ]
    hdr_i = next(
        (
            i
            for i, r in enumerate(rows)
            if any(c.lower().startswith("component") for c in r)
            and any(c.lower() in ("qty", "quantity") for c in r)
        ),
        None,
    )
    if hdr_i is None:
        raise BoqParseError("No header row with Components and Qty was found.")
    hdr = [c.lower() for c in rows[hdr_i]]

    def col(*names: str) -> int | None:
        for i, h in enumerate(hdr):
            if any(h.startswith(n) for n in names):
                return i
        return None

    ci, qi = col("component"), col("qty", "quantity")
    pi, ai, si = col("price", "rate"), col("amount", "total"), col("sr")
    priced = pi is not None
    lines: list[BoqLine] = []
    group = section = ""
    for r in rows[hdr_i + 1 :]:
        comp = r[ci] if ci is not None and ci < len(r) else ""
        qty_s = r[qi] if qi is not None and qi < len(r) else ""
        if comp.lower() in GROUPS and not qty_s:
            group, section = comp.title(), ""
            continue
        if comp and not qty_s:
            section = comp
            continue
        if not comp or not qty_s:
            continue
        try:
            qty = int(Decimal(qty_s.replace(",", "")))
        except (InvalidOperation, ValueError):
            continue
        price = _to_decimal(r[pi]) if priced and pi is not None and pi < len(r) and r[pi] else None
        amount = _to_decimal(r[ai]) if ai is not None and ai < len(r) and r[ai] else None
        issues: list[str] = []
        conf = 1.0
        if price is not None and amount is not None and price * qty != amount:
            issues.append(f"Amount {amount} is not quantity {qty} times price {price}.")
            conf = 0.5
        ref = r[si] if si is not None and si < len(r) else ""
        om = _OPTION.match(ref)
        lines.append(
            BoqLine(
                ref,
                om.group(2) if om else "",
                group,
                section,
                comp,
                "",
                [],
                qty,
                price,
                amount,
                conf,
                issues,
            )
        )
    if not lines:
        raise BoqParseError("No BOQ lines were found below the header.")
    return BoqDocument("quotation" if priced else "summary", None, None, None, None, None, lines)


def detect_xlsx(data: bytes) -> float:
    try:
        parse_xlsx(data)
    except BoqParseError:
        return 0.0
    except Exception:
        return 0.0
    return 0.8


def parsed_at() -> datetime:  # pragma: no cover - trivial
    return datetime.now()
