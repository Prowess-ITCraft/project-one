"""PrismSuiteParserV1: the Word (.docx) report layout of September 2026 (sample PS-10092026-SHA).

How it reads the report:
1. Walk the document body in order, tracking the current section ("▌ Heading") and the last
   short paragraph before each table (its label, e.g. "Firewall" or "1b. License Upgrade ...").
2. Recognise each table by its header row (normalised column names), never by position, so
   extra or reordered tables in later reports do not break it.
3. Fill the AuditSnapshot, recording one FieldReport per important field: ok, missing,
   unreadable, or conflict (when the report states different values in different places).
Nothing here raises for bad content. Only an unreadable file raises ParseFailure.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from decimal import Decimal

from app.modules.prismsuite.parsers import values as v
from app.modules.prismsuite.parsers.base import ParseFailure, ParseResult, register
from app.modules.prismsuite.snapshot import (
    AuditSnapshot,
    AvReplacement,
    ConflictingAv,
    CountItem,
    Device,
    DeviceCategory,
    Endpoint,
    FieldReport,
    FieldStatus,
    LicenceUpgrade,
    OsUpgrade,
    RamUpgrade,
    ReadReport,
    Recommendation,
    Score,
    SecurityComponent,
    SystemRecommendation,
    Vulnerability,
)

PARSER_NAME = "prismsuite.docx.v1"
MAX_TABLES = 500
MAX_ROWS = 5000

KNOWN_AV = (
    "mcafee",
    "quick heal",
    "seqrite",
    "windows defender",
    "microsoft defender",
    "sophos",
    "kaspersky",
    "eset",
    "bitdefender",
    "symantec",
    "norton",
    "trend micro",
    "avast",
    "avg",
    "k7",
    "crowdstrike",
    "sentinelone",
    "acronis",
    "npav",
    "net protector",
)

SECTION_ALIASES = {
    "introduction": "introduction",
    "infrastructure diagram": "diagram",
    "network infrastructure": "network",
    "prismsuite summary report": "summary",
    "os and applications on systems": "os_apps",
    "data network security components": "security_components",
    "gap analysis recommendations": "gap",
    "ceo summary": "ceo",
    "vulnerabilities associated risks": "vulns",
    "health quotient of it infra": "health",
    "high availability grading": "ha",
    "performance grading": "performance",
    "it security risk assessment": "security",
    "server details": "server_details",
    "network gateway security appliance firewall": "firewall_details",
    "backup device": "backup_details",
    "system details with comments": "systems",
    "in depth recommendations per system": "per_system",
}

DEVICE_LABELS = {
    "server": DeviceCategory.SERVER,
    "servers": DeviceCategory.SERVER,
    "firewall": DeviceCategory.FIREWALL,
    "router": DeviceCategory.ROUTER,
    "routers": DeviceCategory.ROUTER,
    "backup": DeviceCategory.NAS,
    "nas": DeviceCategory.NAS,
    "switches": DeviceCategory.SWITCH,
    "switch": DeviceCategory.SWITCH,
    "access point": DeviceCategory.ACCESS_POINT,
    "access points": DeviceCategory.ACCESS_POINT,
}

COMPONENT_KEYS = {
    "server": "server",
    "servers": "server",
    "firewall": "firewall",
    "nas": "nas",
    "nas backup": "nas",
    "backup": "nas",
    "switches": "switch",
    "switch": "switch",
    "router": "router",
    "routers": "router",
    "access point": "access_point",
    "access points": "access_point",
    "endpoints": "endpoints",
}

HEADLINE_LABELS = {
    "overall system health": "system_health",
    "average system health": "system_health",
    "average it structure health": "it_structure_health",
    "performance of infra": "performance",
    "ha of infra": "high_availability",
    "overall it security posture": "security",
}


@dataclass
class Table:
    index: int
    section: str  # normalised section key, or "unknown:<text>"
    label: str  # last paragraph before the table
    header: list[str]  # normalised header cells
    raw_header: list[str]  # first-row cell text as written
    rows: list[list[str]]  # cleaned body cells (header excluded)
    raw_rows: list[list[str]]  # un-normalised cell text, newlines kept
    nested: list[list[Table]] = field(default_factory=list)  # per cell of the first row

    def col(self, *names: str) -> int | None:
        for n in names:
            for i, h in enumerate(self.header):
                if h == n:
                    return i
        for n in names:
            for i, h in enumerate(self.header):
                if n in h:
                    return i
        return None

    def where(self) -> str:
        return f"table {self.index + 1} ('{self.label or self.section}')"


@dataclass
class Doc:
    paragraphs: list[tuple[str, str]]  # (section, text)
    tables: list[Table]
    core_author: str | None
    core_last_modified_by: str | None


def _cell_texts(tbl: object) -> tuple[list[list[str]], list[list[object]]]:
    """Rows of cell text with horizontally merged cells collapsed, plus the cell objects."""
    rows: list[list[str]] = []
    cells_out: list[list[object]] = []
    for row in tbl.rows[:MAX_ROWS]:  # type: ignore[attr-defined]
        texts: list[str] = []
        cells: list[object] = []
        prev = None
        for c in row.cells:
            if c._tc is prev:
                continue
            prev = c._tc
            texts.append(c.text)
            cells.append(c)
        rows.append(texts)
        cells_out.append(cells)
    return rows, cells_out


def _make_table(tbl: object, index: int, section: str, label: str) -> Table:
    raw_rows, cells = _cell_texts(tbl)
    header = [v.norm_key(c) for c in raw_rows[0]] if raw_rows else []
    body_raw = raw_rows[1:]
    nested: list[list[Table]] = []
    if cells:
        for c in cells[0]:
            inner = [_make_table(t, index, section, label) for t in getattr(c, "tables", [])[:20]]
            nested.append(inner)
    return Table(
        index=index,
        section=section,
        label=label,
        header=header,
        raw_header=[v.clean(c) for c in raw_rows[0]] if raw_rows else [],
        rows=[[v.clean(x).replace("\n", " ") for x in r] for r in body_raw],
        raw_rows=[[v.clean(x) for x in r] for r in body_raw],
        nested=nested,
    )


def _section_key(text: str) -> str:
    k = v.norm_key(text)
    for alias, key in SECTION_ALIASES.items():
        if alias in k:
            return key
    return f"unknown:{text[:80]}"


def load(data: bytes) -> Doc:
    try:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise ParseFailure("python-docx is not installed") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ParseFailure("The file is not a readable Word document.") from exc
    except Exception as exc:  # python-docx raises assorted lxml errors on corrupt XML
        raise ParseFailure("The Word document is damaged.") from exc

    section = "header"
    label = ""
    paragraphs: list[tuple[str, str]] = []
    tables: list[Table] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, document)
            text = v.clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if p.text.lstrip().startswith("▌") or style.lower().startswith("heading"):
                section = _section_key(text)
                label = ""
            else:
                label = text if len(text) <= 120 else label
            paragraphs.append((section, text))
        elif child.tag == qn("w:tbl") and len(tables) < MAX_TABLES:
            tables.append(_make_table(DocxTable(child, document), len(tables), section, label))
    props = document.core_properties
    return Doc(
        paragraphs=paragraphs,
        tables=tables,
        core_author=v.text_or_none(props.author),
        core_last_modified_by=v.text_or_none(props.last_modified_by),
    )


class _Reader:
    """Holds the snapshot under construction and the field-level read report."""

    def __init__(self, doc: Doc) -> None:
        self.doc = doc
        self.s = AuditSnapshot()
        self.fields: dict[str, FieldReport] = {}
        self.used_tables: set[int] = set()

    # -- report helpers
    def ok(self, path: str, source: str, *, required: bool = False, raw: str | None = None) -> None:
        if path in self.fields and self.fields[path].status == FieldStatus.CONFLICT:
            return
        self.fields[path] = FieldReport(
            path=path, status=FieldStatus.OK, source=source, required=required, raw=raw
        )

    def missing(self, path: str, message: str, *, required: bool = False) -> None:
        if path not in self.fields:
            self.fields[path] = FieldReport(
                path=path, status=FieldStatus.MISSING, message=message, required=required
            )

    def unreadable(
        self,
        path: str,
        message: str,
        source: str,
        *,
        required: bool = False,
        raw: str | None = None,
    ) -> None:
        self.fields[path] = FieldReport(
            path=path,
            status=FieldStatus.UNREADABLE,
            message=message,
            source=source,
            required=required,
            raw=raw,
        )

    def conflict(self, path: str, message: str, source: str, *, required: bool = False) -> None:
        self.fields[path] = FieldReport(
            path=path,
            status=FieldStatus.CONFLICT,
            message=message,
            source=source,
            required=required,
        )

    def tables_with(self, *headers: str, section: str | None = None) -> Iterator[Table]:
        for t in self.doc.tables:
            if section and t.section != section:
                continue
            if all(any(h == c or h in c for c in t.header) for h in headers):
                yield t

    def use(self, t: Table) -> Table:
        self.used_tables.add(t.index)
        return t

    def paragraph(self, pattern: str) -> tuple[str, re.Match[str]] | None:
        rx = re.compile(pattern, re.I)
        for section, text in self.doc.paragraphs:
            if m := rx.search(text):
                return section, m
        return None

    # ------------------------------------------------------------------ extractors
    def header(self) -> None:
        h = self.s.header
        head = [t for s, t in self.doc.paragraphs if s == "header"]
        if head:
            h.customer_name = head[0][:250]
            self.ok("/header/customer_name", "first line of the report", required=True)
        else:
            self.missing("/header/customer_name", "The report has no title block.", required=True)
        for text in head[1:4]:
            if "report reference" in text.lower() or "prepared" in text.lower():
                continue
            if (d := v.parse_date(text)) is not None:
                h.audit_date = d
                loc = re.split(r"\s[·|•]\s|\s{2,}", text)[0].strip(" ·")
                h.customer_location = loc[:250] or None
                self.ok("/header/audit_date", "title block", required=True, raw=text)
                break
        ref_hit = self.paragraph(r"Report Reference\s*:?\s*(\S+)")
        if ref_hit:
            ref, ref_date = v.parse_report_reference(ref_hit[1].group(1))
            if ref:
                h.report_reference = ref
                self.ok(
                    "/header/report_reference",
                    "title block",
                    required=True,
                    raw=ref_hit[1].group(0),
                )
                if h.audit_date is None and ref_date:
                    h.audit_date = ref_date
                    self.ok("/header/audit_date", "date inside the report reference", required=True)
                elif h.audit_date and ref_date and h.audit_date != ref_date:
                    self.conflict(
                        "/header/audit_date",
                        f"Title shows {h.audit_date}, reference shows {ref_date}.",
                        "title block",
                        required=True,
                    )
            else:
                self.unreadable(
                    "/header/report_reference",
                    "Expected PS-DDMMYYYY-XXX.",
                    "title block",
                    required=True,
                    raw=ref_hit[1].group(1),
                )
        else:
            self.missing(
                "/header/report_reference", "No 'Report Reference' line found.", required=True
            )
        self.missing("/header/audit_date", "No audit date found in the title block.", required=True)
        if prep := self.paragraph(r"Prepared by\s+(.+)"):
            h.prepared_by = prep[1].group(1).strip()[:120]
        auditor = self.paragraph(r"(?:Audited by|Auditor)\s*:\s*(.+)")
        if auditor:
            h.auditor = auditor[1].group(1).strip()[:120]
            h.auditor_source = "report"
            self.ok("/header/auditor", "report text")
        else:
            name = self.doc.core_last_modified_by or self.doc.core_author
            if name and name.lower() not in ("python-docx", "prismsuite", "microsoft office user"):
                h.auditor = name[:120]
                h.auditor_source = "document_properties"
                self.fields["/header/auditor"] = FieldReport(
                    path="/header/auditor",
                    status=FieldStatus.OK,
                    message="Taken from the Word file's 'last modified by' property. Confirm during review.",
                    source="document properties",
                )
            else:
                self.missing("/header/auditor", "The report does not name the auditor.")

    def devices(self) -> None:
        found: set[DeviceCategory] = set()
        for t in self.tables_with("brand model", "type", "configuration"):
            cat = DEVICE_LABELS.get(v.norm_key(t.label))
            if cat is None:
                continue
            self.use(t)
            bi, ti, ai, ci = (
                t.col("brand model"),
                t.col("type"),
                t.col("age"),
                t.col("configuration"),
            )
            for row, raw in zip(t.rows, t.raw_rows, strict=False):
                brand_model = v.text_or_none(row[bi]) if bi is not None and bi < len(row) else None
                if not brand_model:
                    continue
                dtype = v.text_or_none(row[ti]) if ti is not None and ti < len(row) else None
                config = v.text_or_none(raw[ci]) if ci is not None and ci < len(raw) else None
                d = Device(
                    category=cat,
                    brand_model=brand_model[:200],
                    brand=brand_model.split()[0][:60],
                    model=" ".join(brand_model.split()[1:])[:120] or None,
                    device_type=dtype,
                    age=v.text_or_none(row[ai]) if ai is not None and ai < len(row) else None,
                    configuration=config,
                )
                if cat == DeviceCategory.SWITCH and dtype:
                    low = dtype.lower()
                    d.managed = (
                        False if "unmanaged" in low else (True if "managed" in low else None)
                    )
                if config:
                    if m := re.search(r"Ports?\s*:\s*(\d+)", config, re.I):
                        d.ports = int(m.group(1))
                    if m := re.search(r"\((\d{1,3}(?:\.\d+)?)\s*%", config):
                        d.storage_used_percent = Decimal(m.group(1))
                    if m := re.search(
                        r"(SonicOS|FortiOS|SFOS|DSM|firmware|FW)\s*:?\s*([\w.\-]+(?:\s+Update\s+\d+)?)",
                        config,
                        re.I,
                    ):
                        d.firmware = (
                            f"{m.group(1)} {m.group(2)}"
                            if m.group(1).lower() not in ("fw", "firmware")
                            else m.group(2)
                        )
                self.s.devices.append(d)
                found.add(cat)
        for cat in (
            DeviceCategory.SERVER,
            DeviceCategory.FIREWALL,
            DeviceCategory.NAS,
            DeviceCategory.SWITCH,
            DeviceCategory.ROUTER,
        ):
            path = f"/devices[{cat.value}]"
            if cat in found:
                self.ok(path, "Network Infrastructure tables")
            else:
                self.missing(
                    path, f"No {cat.value.replace('_', ' ')} table in Network Infrastructure."
                )

    def asset_summary(self) -> None:
        a = self.s.assets
        mapping = {
            "laptops": "laptops",
            "laptop": "laptops",
            "desktops": "desktops",
            "desktop": "desktops",
            "server": "servers",
            "servers": "servers",
            "firewall": "firewalls",
            "firewalls": "firewalls",
            "backup device": "backup_devices",
            "backup devices": "backup_devices",
            "switches": "switches",
            "switch": "switches",
            "routers": "routers",
            "router": "routers",
            "access point": "access_points",
            "access points": "access_points",
        }
        tables = list(self.tables_with("system type", "quantity"))
        if not tables:
            self.missing("/assets", "Asset Summary table not found.", required=True)
        for t in tables[:1]:
            self.use(t)
            si, qi = t.col("system type"), t.col("quantity")
            for row in t.rows:
                if si is None or qi is None or max(si, qi) >= len(row):
                    continue
                key = mapping.get(v.norm_key(row[si]))
                if key is None:
                    continue
                n = v.parse_int(row[qi])
                if n is None and not v.is_na(row[qi]):
                    self.unreadable(
                        f"/assets/{key}", "Quantity is not a number.", t.where(), raw=row[qi]
                    )
                    continue
                setattr(a, key, n)
                self.ok(f"/assets/{key}", t.where())
        if a.laptops is not None or a.desktops is not None:
            a.endpoints_total = (a.laptops or 0) + (a.desktops or 0)
            self.ok("/assets/endpoints_total", "Asset Summary (laptops + desktops)", required=True)
        else:
            self.missing(
                "/assets/endpoints_total",
                "Laptop and desktop counts were not found.",
                required=True,
            )
        hit = self.paragraph(r"Cumulative Storage\s*:\s*(.+)")
        if hit:
            text = hit[1].group(1)
            for label, attr in (
                ("total", "storage_total_tb"),
                ("used", "storage_used_tb"),
                ("free", "storage_free_tb"),
            ):
                if m := re.search(rf"([\d.,]+)\s*TB\s*{label}", text, re.I):
                    setattr(a, attr, v.to_decimal(m.group(1)))

    def os_and_software(self) -> None:
        for t in self.tables_with("operating system", "systems"):
            if len(t.header) != 2:
                continue
            self.use(t)
            for row in t.rows:
                if len(row) >= 2 and not v.is_na(row[0]):
                    self.s.os_distribution.append(
                        CountItem(name=row[0][:120], count=v.parse_int(row[1]))
                    )
            self.ok("/os_distribution", t.where())
            break
        else:
            self.missing("/os_distribution", "OS Distribution table not found.")
        for t in self.tables_with("name", "software type", "systems"):
            self.use(t)
            ni, ci, ki = t.col("name"), t.col("software type"), t.col("systems")
            for row in t.rows:
                if ni is None or ki is None or max(ni, ki) >= len(row) or v.is_na(row[ni]):
                    continue
                self.s.installed_software.append(
                    CountItem(
                        name=row[ni][:120],
                        count=v.parse_int(row[ki]),
                        category=row[ci][:80] if ci is not None and ci < len(row) else None,
                    )
                )
            self.ok("/installed_software", t.where())
            break

    def security_components(self) -> None:
        for t in self.doc.tables:
            if t.section != "security_components" or not t.nested:
                continue
            scopes = ["server" if "server" in h else "endpoint" for h in t.header]
            got = False
            for i, inner_tables in enumerate(t.nested):
                scope = scopes[i] if i < len(scopes) else "endpoint"
                for inner in inner_tables:
                    ti, pi, qi = (
                        inner.col("technology"),
                        inner.col("product"),
                        inner.col("quantity"),
                    )
                    if ti is None:
                        continue
                    for row in inner.rows:
                        if ti >= len(row) or v.is_na(row[ti]):
                            continue
                        product = row[pi] if pi is not None and pi < len(row) else ""
                        detected = not ("not detected" in product.lower() or v.is_na(product))
                        self.s.security_components.append(
                            SecurityComponent(
                                scope=scope,
                                technology=row[ti][:60],
                                product=v.clean(product).strip("●• ")[:200] if detected else None,
                                detected=detected,
                                quantity=v.parse_int(row[qi])
                                if qi is not None and qi < len(row)
                                else None,
                            )
                        )
                        got = True
            if got:
                self.use(t)
                self.ok("/security_components", t.where())
                return
        self.missing("/security_components", "Security components tables not found.")

    def _set_headline(self, key: str, score: Score, source: str) -> None:
        current: Score = getattr(self.s.scores, key)
        path = f"/scores/{key}/value"
        required = key in ("security", "high_availability", "performance")
        if score.value is None:
            if current.value is None:
                self.unreadable(
                    path, "Score could not be read.", source, required=required, raw=score.raw
                )
            return
        if current.value is not None and current.value != score.value:
            self.conflict(
                path,
                f"{current.value} in one place, {score.value} in {source}.",
                source,
                required=required,
            )
            return
        merged = current.model_copy(
            update={k: val for k, val in score.model_dump().items() if val is not None}
        )
        setattr(self.s.scores, key, merged)
        self.ok(path, source, required=required, raw=score.raw)

    def headline_scores(self) -> None:
        # One-row summary cards: ['', 'Overall System Health\nComment: ...', '74.2%\nNEEDS ATTENTION']
        for t in self.doc.tables:
            if t.header and len(t.header) == 3 and not t.rows:
                name = t.header[1]
                key = next(
                    (k for label, k in HEADLINE_LABELS.items() if name.startswith(label)), None
                )
                if key is None:
                    continue
                self.use(t)
                text = t.raw_header[2]
                score = v.parse_score(text.split("\n")[0])
                if score.value is not None:
                    score.label = (text.split("\n")[1] if "\n" in text else None) or score.label
                self._set_headline(key, score, t.where())
        # CEO summary: Name | Score | Comment | Ideal Range
        for t in self.tables_with("name", "score", "ideal range"):
            self.use(t)
            ni, si, ri = t.col("name"), t.col("score"), t.col("ideal range")
            for row in t.rows:
                if ni is None or si is None or max(ni, si) >= len(row):
                    continue
                key = HEADLINE_LABELS.get(v.norm_key(row[ni]))
                if key is None:
                    continue
                score = v.parse_score(row[si])
                if ri is not None and ri < len(row):
                    score.ideal_range = v.text_or_none(row[ri])
                self._set_headline(key, score, t.where())
        for key in (
            "system_health",
            "it_structure_health",
            "performance",
            "high_availability",
            "security",
        ):
            self.missing(
                f"/scores/{key}/value",
                "Score not found.",
                required=key in ("security", "high_availability", "performance"),
            )

    def component_scores(self) -> None:
        specs: list[tuple[str, str, str, Callable[[Score], None] | None]] = [
            ("health", "device", "score", None),
            ("high_availability", "component", "ha score", None),
            ("performance", "component", "performance score", None),
            ("security", "security component", "risk score", None),
        ]
        for attr, name_col, score_col, _ in specs:
            target: dict[str, Score] = getattr(self.s.component_scores, attr)
            tables = [
                t for t in self.tables_with(name_col, score_col) if t.col(score_col) is not None
            ]
            if attr == "health":
                tables = [
                    t
                    for t in tables
                    if t.col("ha score") is None
                    and t.col("performance score") is None
                    and t.col("risk score") is None
                ]
            for t in tables:
                self.use(t)
                ni, si = t.col(name_col), t.col(score_col)
                assert ni is not None and si is not None
                for row in t.rows:
                    if max(ni, si) >= len(row):
                        continue
                    name = v.norm_key(row[ni])
                    score = v.parse_score(row[si])
                    if name.startswith("overall"):
                        self._check_overall(attr, score, t.where())
                        continue
                    key = COMPONENT_KEYS.get(name)
                    if key is None:
                        continue
                    if score.value is None and not v.is_na(row[si]):
                        self.unreadable(
                            f"/component_scores/{attr}/{key}",
                            "Score could not be read.",
                            t.where(),
                            raw=row[si],
                        )
                    target[key] = score
            if target:
                self.ok(f"/component_scores/{attr}", "grading tables")
            else:
                self.missing(
                    f"/component_scores/{attr}", f"No {attr.replace('_', ' ')} grading table found."
                )
        # "Overall Security Score: 58.7 / 100 - AT RISK" style lines
        for pattern, key in (
            (r"Overall Security Score\s*:\s*(.+)", "security"),
            (r"Overall HA Score\s*:\s*(.+)", "high_availability"),
            (r"Overall Performance Score\s*:\s*(.+)", "performance"),
            (r"Overall Endpoint Health Score\s*:\s*(.+)", "system_health"),
            (r"Overall Health Score\s*:\s*(.+)", "it_structure_health"),
        ):
            if hit := self.paragraph(pattern):
                self._set_headline(key, v.parse_score(hit[1].group(1)), f"'{hit[1].group(0)[:60]}'")

    def _check_overall(self, attr: str, score: Score, source: str) -> None:
        key = {
            "health": "it_structure_health",
            "high_availability": "high_availability",
            "performance": "performance",
            "security": "security",
        }[attr]
        self._set_headline(key, score, source)

    def upgrades(self) -> None:
        u = self.s.upgrades
        for t in self.tables_with("upgrade type", "count"):
            self.use(t)
            ui, ci = t.col("upgrade type"), t.col("count")
            assert ui is not None and ci is not None
            for row in t.rows:
                if max(ui, ci) >= len(row) or v.is_na(row[ui]):
                    continue
                count = v.parse_int(row[ci])
                if count is None:
                    self.unreadable(
                        "/upgrades/ram", "Count is not a number.", t.where(), raw=row[ci]
                    )
                    continue
                cur = re.search(r"current\s*:\s*([\d.]+)\s*GB", row[ui], re.I)
                req = re.search(r"required\s*:\s*([\d.]+)\s*GB", row[ui], re.I)
                u.ram.append(
                    RamUpgrade(
                        description=row[ui][:300],
                        count=count,
                        current_gb=v.to_decimal(cur.group(1)) if cur else None,
                        required_gb=v.to_decimal(req.group(1)) if req else None,
                    )
                )
            self.ok("/upgrades/ram", t.where())
            break
        else:
            self.missing("/upgrades/ram", "Hardware Upgrade Recommendation table not found.")

        for t in self.tables_with("current version", "upgrade to", "count"):
            self.use(t)
            a, b, c = t.col("current version"), t.col("upgrade to"), t.col("count")
            assert a is not None and b is not None and c is not None
            for row in t.rows:
                if max(a, b, c) >= len(row) or v.is_na(row[a]):
                    continue
                n = v.parse_int(row[c])
                if n is not None:
                    u.licence.append(
                        LicenceUpgrade(
                            current=row[a][:120], upgrade_to=v.text_or_none(row[b]), count=n
                        )
                    )
            self.ok("/upgrades/licence", t.where())
            break
        else:
            self.missing("/upgrades/licence", "License Upgrade Recommendation table not found.")

        for t in self.tables_with("current os", "upgrade to", "count"):
            self.use(t)
            a, b, c = t.col("current os"), t.col("upgrade to"), t.col("count")
            assert a is not None and b is not None and c is not None
            for row in t.rows:
                if max(a, b, c) >= len(row) or v.is_na(row[a]):
                    continue
                n = v.parse_int(row[c])
                if n is not None:
                    u.os.append(
                        OsUpgrade(current=row[a][:120], upgrade_to=v.text_or_none(row[b]), count=n)
                    )
            self.ok("/upgrades/os", t.where())
            break
        else:
            self.missing("/upgrades/os", "OS Upgrade Recommendation table not found.")

        for t in self.tables_with("current av", "count"):
            self.use(t)
            a, c, b = t.col("current av"), t.col("count"), t.col("upgrade to")
            assert a is not None and c is not None
            for row in t.rows:
                if max(a, c) >= len(row) or v.is_na(row[a]):
                    continue
                n = v.parse_int(row[c])
                if n is not None:
                    u.endpoint_protection.append(
                        AvReplacement(
                            current_av=row[a][:120],
                            count=n,
                            upgrade_to=v.text_or_none(row[b])
                            if b is not None and b < len(row)
                            else None,
                        )
                    )
            self.ok("/upgrades/endpoint_protection", t.where())
            break
        else:
            self.missing(
                "/upgrades/endpoint_protection",
                "Endpoint Protection Recommendation table not found.",
            )

        for t in self.tables_with("system", "products installed"):
            self.use(t)
            si, pi, ai = t.col("system"), t.col("products installed"), t.col("action")
            assert si is not None and pi is not None
            for row in t.rows:
                if max(si, pi) >= len(row) or v.is_na(row[si]):
                    continue
                u.conflicting_av.append(
                    ConflictingAv(
                        system=row[si][:200],
                        products=v.split_list(row[pi], commas=True),
                        action=v.text_or_none(row[ai])
                        if ai is not None and ai < len(row)
                        else None,
                    )
                )
            self.ok("/upgrades/conflicting_av", t.where())
            break
        else:
            self.missing(
                "/upgrades/conflicting_av", "Conflicting Antivirus Cleanup table not found."
            )

    def vulnerabilities(self) -> None:
        hit = self.paragraph(r"Total Vulnerabilities Identified\s*:?\s*([\d,]+)")
        if hit:
            self.s.vulnerabilities.total = v.parse_int(hit[1].group(1))
            self.ok(
                "/vulnerabilities/total",
                "Vulnerabilities section",
                required=True,
                raw=hit[1].group(0),
            )
        else:
            self.missing(
                "/vulnerabilities/total",
                "No 'Total Vulnerabilities Identified' line.",
                required=True,
            )
        for t in self.tables_with("cve id", "vulnerability"):
            self.use(t)
            pi, ci, vi, ri, mi = (
                t.col("priority"),
                t.col("cve id"),
                t.col("vulnerability"),
                t.col("risk factor"),
                t.col("recommendation"),
            )
            for row in t.rows:
                if vi is None or vi >= len(row) or v.is_na(row[vi]):
                    continue
                text = row[vi]
                title, _, desc = text.partition(":")
                sev, cvss = (
                    v.parse_severity(row[ri]) if ri is not None and ri < len(row) else (None, None)
                )
                cve = v.find_cve(row[ci]) if ci is not None and ci < len(row) else None
                self.s.vulnerabilities.listed.append(
                    Vulnerability(
                        priority=v.parse_int(row[pi]) if pi is not None and pi < len(row) else None,
                        cve_id=cve or v.find_cve(text),
                        title=title.strip()[:300],
                        description=desc.strip()[:2000] or None,
                        severity=sev,
                        cvss=cvss,
                        recommendation=(v.text_or_none(row[mi]) or "")[:2000] or None
                        if mi is not None and mi < len(row)
                        else None,
                    )
                )
            missing_cve = sum(1 for x in self.s.vulnerabilities.listed if x.cve_id is None)
            self.ok("/vulnerabilities/listed", t.where())
            if missing_cve:
                self.fields[
                    "/vulnerabilities/listed"
                ].message = f"{missing_cve} listed vulnerability without a CVE id."
            break
        else:
            self.missing("/vulnerabilities/listed", "Top vulnerabilities table not found.")

    def recommendations(self) -> None:
        for t in self.tables_with("it infra", "recommendation"):
            self.use(t)
            ai, ri, di = t.col("it infra"), t.col("recommendation"), t.col("description")
            assert ai is not None and ri is not None
            for row in t.rows:
                if max(ai, ri) >= len(row) or v.is_na(row[ri]):
                    continue
                self.s.recommendations.append(
                    Recommendation(
                        area=row[ai][:80],
                        recommendation=row[ri][:500],
                        description=v.text_or_none(row[di])
                        if di is not None and di < len(row)
                        else None,
                    )
                )
            self.ok("/recommendations", t.where())
            break
        else:
            self.missing("/recommendations", "Product recommendations table not found.")
        for section, text in self.doc.paragraphs:
            if section == "vulns" and re.match(
                r"^(sanitize|sanitise|note|important)\b", text, re.I
            ):
                self.s.notes.append(text[:500])

    def detail_tables(self) -> None:
        for section, cat in (
            ("server_details", DeviceCategory.SERVER),
            ("firewall_details", DeviceCategory.FIREWALL),
            ("backup_details", DeviceCategory.NAS),
        ):
            for t in self.tables_with("parameter", "value", section=section):
                self.use(t)
                details = {
                    row[0][:120]: v.text_or_none(row[1])
                    for row in t.rows
                    if len(row) >= 2 and row[0]
                }
                device = self._match_device(cat, t.label)
                if device is None:
                    device = Device(
                        category=cat, brand_model=t.label[:200] or cat.value, details={}
                    )
                    self.s.devices.append(device)
                device.details.update(details)
                fw = next(
                    (
                        val
                        for k, val in details.items()
                        if "firmware" in k.lower() and "version" in k.lower()
                    ),
                    None,
                )
                if fw:
                    if (
                        device.firmware
                        and v.norm_key(device.firmware) not in v.norm_key(fw)
                        and v.norm_key(fw) not in v.norm_key(device.firmware)
                    ):
                        self.conflict(
                            f"/devices[{cat.value}]/firmware",
                            f"Firmware shown as '{device.firmware}' and '{fw}'.",
                            t.where(),
                        )
                    device.firmware = fw[:120]
                if cat == DeviceCategory.SERVER and (hard := details.get("Server Hardening")):
                    score = v.parse_score(hard, default_out_of=Decimal(10))
                    if score.value is not None:
                        self.s.server_hardening_score = score
                        self.ok("/server_hardening_score", t.where(), raw=hard)
                    else:
                        self.unreadable(
                            "/server_hardening_score",
                            "Benchmark score could not be read.",
                            t.where(),
                            raw=hard,
                        )
                if cat == DeviceCategory.NAS and (used := details.get("Drive Utilised")):
                    if m := re.search(r"(\d{1,3}(?:\.\d+)?)\s*%", used):
                        device.storage_used_percent = Decimal(m.group(1))
                if cat == DeviceCategory.FIREWALL:
                    self._firewall_ha_check(device, t.where())
        if self.s.server_hardening_score is None:
            self.missing("/server_hardening_score", "Server hardening benchmark not found.")

    def _match_device(self, cat: DeviceCategory, label: str) -> Device | None:
        from rapidfuzz import fuzz

        candidates = self.s.devices_of(cat)
        if not candidates:
            return None
        if len(candidates) == 1:
            return candidates[0]
        best = max(candidates, key=lambda d: fuzz.token_set_ratio(d.brand_model, label))
        return best

    def _firewall_ha_check(self, fw: Device, source: str) -> None:
        statements: list[str] = []
        for text in filter(None, [fw.configuration, fw.details.get("High Availability")]):
            statements.append(text)
        for t in self.tables_with("it infra", "current config"):
            self.use(t)
            for row in t.rows:
                if len(row) >= 3 and v.norm_key(row[1]) == "firewall":
                    statements.append(row[2])
        for text in statements:
            m = re.search(r"FW\s*:\s*([\w.\-]+)", text)
            if m and fw.firmware and not _same_version(m.group(1), fw.firmware):
                self.conflict(
                    "/devices[firewall]/firmware",
                    f"Firmware shown as '{fw.firmware}' and as '{m.group(1)}' in Infra Setup.",
                    source,
                )
        says: set[bool] = set()
        for text in statements:
            low = text.lower()
            if re.search(
                r"ha\s*:?\s*not configured|high availability\s*:?\s*not configured|^not configured",
                low,
            ):
                says.add(False)
            elif re.search(r"ha\s*:\s*configured with|ha\s*:\s*configured\b", low):
                says.add(True)
        if len(says) > 1:
            self.conflict(
                "/devices[firewall]/high_availability",
                "The report says both that firewall HA is configured and that it is not.",
                source,
            )

    def endpoints(self) -> None:
        tables = list(self.tables_with("system name", "hardware specs"))
        if not tables:
            self.missing("/endpoints", "System Details table not found.")
        for t in tables[:1]:
            self.use(t)
            c = {
                k: t.col(k)
                for k in (
                    "user name",
                    "system name",
                    "type",
                    "hardware specs",
                    "operating system",
                    "purchase date",
                    "approved software",
                    "comments",
                )
            }
            for row, raw in zip(t.rows, t.raw_rows, strict=False):

                def get(
                    k: str,
                    *,
                    keep_lines: bool = False,
                    _row: list[str] = row,
                    _raw: list[str] = raw,
                ) -> str | None:
                    i = c[k]
                    if i is None or i >= len(_row):
                        return None
                    return v.text_or_none(_raw[i] if keep_lines else _row[i])

                name = get("system name")
                if not name:
                    continue
                specs = get("hardware specs", keep_lines=True) or ""
                software = v.split_list(get("approved software", keep_lines=True))
                ram = re.search(r"RAM\s*:\s*([\d.]+)\s*GB", specs, re.I)
                self.s.endpoints.append(
                    Endpoint(
                        user_name=get("user name"),
                        system_name=name[:120],
                        system_type=get("type"),
                        processor=_line_value(specs, "Processor"),
                        motherboard=_line_value(specs, "Motherboard"),
                        ram_gb=v.to_decimal(ram.group(1)) if ram else None,
                        drives=_line_value(specs, "Drives"),
                        os=get("operating system"),
                        purchase_date=get("purchase date"),
                        software=software,
                        antivirus=[s for s in software if any(av in s.lower() for av in KNOWN_AV)],
                        comments=get("comments"),
                    )
                )
            self.ok("/endpoints", t.where())
        for t in self.tables_with("system", "ram", "storage", "office"):
            self.use(t)
            c2 = {k: t.col(k) for k in ("system", "ram", "storage", "os", "office")}
            for row in t.rows:
                si = c2["system"]
                if si is None or si >= len(row) or v.is_na(row[si]):
                    continue

                def cell(k: str, _row: list[str] = row) -> tuple[str | None, bool | None]:
                    i = c2[k]
                    if i is None or i >= len(_row):
                        return None, None
                    text = v.text_or_none(_row[i])
                    if text is None:
                        return None, None
                    low = text.lower()
                    need = (
                        True
                        if low.startswith("upgrade required")
                        else (False if low.startswith("no upgrade") else None)
                    )
                    return text[:300], need

                ram_t, ram_n = cell("ram")
                st_t, st_n = cell("storage")
                os_t, os_n = cell("os")
                of_t, of_n = cell("office")
                self.s.system_recommendations.append(
                    SystemRecommendation(
                        system_key=row[si][:160],
                        ram=ram_t,
                        ram_upgrade_needed=ram_n,
                        storage=st_t,
                        storage_upgrade_needed=st_n,
                        os=os_t,
                        os_upgrade_needed=os_n,
                        office=of_t,
                        office_upgrade_needed=of_n,
                    )
                )
            self.ok("/system_recommendations", t.where())
            break

    def cross_checks(self) -> None:
        total = self.s.assets.endpoints_total
        if total is None:
            return
        checks = []
        if self.s.endpoints:
            checks.append(("System Details rows", len(self.s.endpoints)))
        os_sum = sum(i.count or 0 for i in self.s.os_distribution)
        if os_sum:
            checks.append(("OS Distribution total", os_sum))
        bad = [f"{name} = {n}" for name, n in checks if n != total]
        if bad:
            self.conflict(
                "/assets/endpoints_total",
                f"Asset Summary gives {total} endpoints but " + ", ".join(bad) + ".",
                "cross-check",
                required=True,
            )
        ram_rows = sum(1 for r in self.s.system_recommendations if r.ram_upgrade_needed)
        if (
            self.s.system_recommendations
            and self.s.upgrades.ram
            and ram_rows != self.s.upgrades.ram_upgrade_count
        ):
            self.conflict(
                "/upgrades/ram",
                f"Gap analysis counts {self.s.upgrades.ram_upgrade_count} RAM upgrades; per-system table flags {ram_rows}.",
                "cross-check",
            )
        office_rows = sum(1 for r in self.s.system_recommendations if r.office_upgrade_needed)
        if (
            self.s.system_recommendations
            and self.s.upgrades.licence
            and office_rows != self.s.upgrades.licence_upgrade_count
        ):
            self.conflict(
                "/upgrades/licence",
                f"Gap analysis counts {self.s.upgrades.licence_upgrade_count} licence upgrades; per-system table flags {office_rows}.",
                "cross-check",
            )

    def run(self) -> ParseResult:
        steps: list[Callable[[], None]] = [
            self.header,
            self.devices,
            self.asset_summary,
            self.os_and_software,
            self.security_components,
            self.headline_scores,
            self.component_scores,
            self.upgrades,
            self.vulnerabilities,
            self.recommendations,
            self.detail_tables,
            self.endpoints,
            self.cross_checks,
        ]
        for step in steps:
            try:
                step()
            except Exception as exc:  # one broken section must not lose the rest
                self.fields[f"/_step/{step.__name__}"] = FieldReport(
                    path=f"/_step/{step.__name__}",
                    status=FieldStatus.UNREADABLE,
                    message=f"This part of the report could not be read ({type(exc).__name__}).",
                )
        unknown_sections = sorted(
            {s.removeprefix("unknown:") for s, _ in self.doc.paragraphs if s.startswith("unknown:")}
        )
        unknown_tables = [
            f"{t.where()} with columns {', '.join(t.header[:6])}"
            for t in self.doc.tables
            if t.index not in self.used_tables and t.header and any(t.rows) and not _is_ignorable(t)
        ]
        report = ReadReport(
            parser=PARSER_NAME,
            fields=sorted(self.fields.values(), key=lambda f: f.path),
            unknown_sections=unknown_sections,
            unknown_tables=unknown_tables[:100],
        )
        return ParseResult(snapshot=self.s, report=report)


def _same_version(a: str, b: str) -> bool:
    """'7.0.1-5145-R5125' vs 'SonicOS 7.0.1-5145-R5125' match; '7.0.1.3' does not."""
    na = re.sub(r"[^0-9]+", ".", a).strip(".")
    nb = re.sub(r"[^0-9]+", ".", b).strip(".")
    return na == nb or na in nb or (nb in na and len(nb) > 3)


def _is_ignorable(t: Table) -> bool:
    """Tables we know about but deliberately do not import."""
    h = set(t.header)
    return {"sr no", "quantity", "pop", "type"} <= h or ({"system", "os", "endpoint security"} <= h)


def _line_value(text: str, key: str) -> str | None:
    m = re.search(rf"{key}\s*:\s*(.+)", text, re.I)
    return (v.text_or_none(m.group(1)) or "")[:200] or None if m else None


class PrismSuiteParserV1:
    name = PARSER_NAME
    kinds = frozenset({"docx"})

    def detect(self, data: bytes) -> float:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                xml = zf.read("word/document.xml")[:400_000]
        except Exception:
            return 0.0
        score = 0.0
        if b"PrismSuite" in xml:
            score += 0.5
        if re.search(rb"PS-\d{8}-[A-Z0-9]{2,6}", xml):
            score += 0.3
        if b"Brand &amp; Model" in xml or b"Asset Summary" in xml:
            score += 0.2
        return min(score, 1.0)

    def parse(self, data: bytes) -> ParseResult:
        doc = load(data)
        return _Reader(doc).run()


register(PrismSuiteParserV1())
