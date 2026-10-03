"""Quotation and summary BOQ as HTML, PDF and Excel, in the ITCraft format.

The layout follows the sample: letterhead, QUOTATION, To block with date and quote reference,
a table of Sr N, Components, Qty, Price, Amount with priority and section rows, bold titles with
bullet inclusions, alternative options numbered 6A and 6B, terms, and a signature block. Money
uses Indian grouping (INR 1,09,653.00). Totals rows appear only when the quote asks for them.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Any

from app.core.documents import RenderedDocument, template_env
from app.core.documents import render_pdf as _render_pdf
from app.core.money import format_inr
from app.core.timeutil import format_long_date, today_ist
from app.modules.boq import draft as d

_env = template_env(Path(__file__).parent / "html")


def build_rows(draft: d.Draft, *, priced: bool) -> tuple[list[dict[str, Any]], str | None]:
    """Flatten the draft into display rows: group, section and line rows, then totals."""
    comp = d.compute(draft, today_ist())
    result = {r.id: r for r in comp.lines}
    rows: list[dict[str, Any]] = []
    last_group = last_section = None
    for n in d.numbered(draft):
        sec = next(
            s for s in draft.sections if s.title == n.section_title and s.group_id == n.group_id
        )
        if n.group_id != last_group:
            rows.append(
                {
                    "kind": "group",
                    "title": next(g.title for g in draft.groups if g.id == n.group_id),
                }
            )
            last_group, last_section = n.group_id, None
        if sec.id != last_section:
            rows.append({"kind": "section", "title": sec.title})
            last_section = sec.id
        ln, r = n.line, result[n.line.id]
        rows.append(
            {
                "kind": "line",
                "ref": n.ref,
                "title": ln.title,
                "description": ln.description,
                "inclusions": ln.inclusions,
                "qty": ln.qty,
                "unit_price": format_inr(ln.unit_price) if ln.unit_price is not None else "",
                "amount": format_inr(r.amount) if r.amount is not None else "",
            }
        )
    note = None
    t = comp.totals
    if priced and (draft.settings.show_totals or draft.settings.show_gst_rows):
        if t.options and not t.complete:
            note = "Alternative options are not added together. The totals below show the lowest and highest choice."

        def rng(a: Any, b: Any) -> str:
            return format_inr(a) if a == b else f"{format_inr(a)} to {format_inr(b)}"

        if draft.settings.show_totals:
            rows.append(
                {
                    "kind": "total",
                    "label": "Sub total",
                    "value": rng(t.subtotal_min, t.subtotal_max),
                }
            )
        if draft.settings.show_gst_rows:
            rows.append(
                {
                    "kind": "total",
                    "label": f"GST @ {str(draft.settings.gst_default).rstrip('0').rstrip('.')}%",
                    "value": rng(t.gst_min, t.gst_max),
                }
            )
        if draft.settings.show_totals:
            rows.append(
                {"kind": "total", "label": "Grand total", "value": rng(t.total_min, t.total_max)}
            )
    return rows, note


def render_html(
    draft: d.Draft,
    company: dict[str, Any],
    *,
    quote_ref: str | None,
    kind: str = "quotation",
    on: date | None = None,
) -> str:
    priced = kind == "quotation"
    rows, note = build_rows(draft, priced=priced)
    s = draft.settings
    ctx = {
        "title": "Quotation" if priced else "BOQ",
        "heading": "QUOTATION" if priced else "BOQ",
        "company": company,
        "customer": s.customer.model_dump(),
        "quote_ref": quote_ref,
        "date_long": format_long_date(on or s.quote_date or today_ist()),
        "intro": s.intro,
        "rows": rows,
        "priced": priced,
        "terms": draft.terms if priced else [],
        "option_note": note,
        "signatory_name": s.signatory_name,
        "signatory_designation": s.signatory_designation,
    }
    return _env.get_template("quotation.html").render(**ctx)


def render_pdf(html: str) -> RenderedDocument:
    """Through the shared WeasyPrint renderer (ADR 0016)."""
    return _render_pdf(html)


def render_xlsx(
    draft: d.Draft, company: dict[str, Any], *, quote_ref: str | None, kind: str = "quotation"
) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    priced = kind == "quotation"
    rows, _ = build_rows(draft, priced=priced)
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Quotation" if priced else "BOQ"
    bold = Font(bold=True)
    ws.append([company["brand"]])
    ws["A1"].font = Font(bold=True, size=16)
    ws.append([company["legal_name"], "", "", "", quote_ref or "Draft"])
    ws.append(
        [
            f"To: M/S {draft.settings.customer.name}",
            "",
            "",
            "",
            format_long_date(draft.settings.quote_date or today_ist()),
        ]
    )
    ws.append([])
    head = ["Sr N", "Components", "Qty"] + (["Price", "Amount"] if priced else [])
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, PatternFill("solid", fgColor="DDDDDD")
    for r in rows:
        if r["kind"] in ("group", "section"):
            ws.append([r["title"]])
            ws[ws.max_row][0].font = bold
        elif r["kind"] == "line":
            text = (
                r["title"]
                + ("\n" + r["description"] if r["description"] else "")
                + "".join(f"\n• {i}" for i in r["inclusions"])
            )
            ws.append(
                [r["ref"], text, r["qty"]] + ([r["unit_price"], r["amount"]] if priced else [])
            )
            ws[ws.max_row][1].alignment = Alignment(wrap_text=True, vertical="top")
        elif r["kind"] == "total":
            ws.append(["", r["label"], "", "", r["value"]])
            ws[ws.max_row][1].font = bold
    if priced and draft.terms:
        ws.append([])
        ws.append(["Terms & Conditions"])
        ws[ws.max_row][0].font = bold
        for t in draft.terms:
            ws.append(["•", t])
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 22
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
